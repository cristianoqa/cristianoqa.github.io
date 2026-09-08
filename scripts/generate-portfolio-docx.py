# -*- coding: utf-8 -*-
"""Generate Flow Home Apps portfolio report (DOCX) for ads + product overview."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pathlib import Path

OUT = Path(r"D:\cristianoqa.github.io\docs\FlowHome-Informe-Apps-Anuncios.docx")
OUT.parent.mkdir(parents=True, exist_ok=True)

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)

def h1(t):
    p = doc.add_heading(t, level=1)
    return p

def h2(t):
    doc.add_heading(t, level=2)

def h3(t):
    doc.add_heading(t, level=3)

def p(t, bold=False):
    para = doc.add_paragraph()
    run = para.add_run(t)
    run.bold = bold
    return para

def bullets(items):
    for i in items:
        doc.add_paragraph(i, style="List Bullet")

# Cover
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("Flow Home Apps")
r.bold = True
r.font.size = Pt(28)

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sr = sub.add_run("Informe de producto para anuncios y publicación")
sr.font.size = Pt(14)

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.add_run("8 septiembre 2026 · Cristiano · desarrollador independiente\n")
meta.add_run("Landing: https://cristianoqa.github.io/")

doc.add_paragraph()
p("Este documento describe cada app: qué es, a quién sirve, funcionalidades clave, importancia comercial y canales de distribución actuales. Sirve como base para creatividades de anuncio, fichas de tienda y comunicación a testers/amigos.")

h1("1. Resumen del portfolio")
bullets([
    "Misiva — WhatsApp programado + IA (solo Android).",
    "MyPass — Gestor de contraseñas local-first (Android + TestFlight iOS).",
    "ReformaPRO — Presupuestos para reformas (PWA + Android + TestFlight iOS).",
    "Monexa — Gastos compartidos en grupo (Android + TestFlight iOS).",
    "Lunera — Salud menstrual / ciclo (Android + PWA + TestFlight iOS).",
    "Miravista — Espejo de pantalla a TV por Wi‑Fi (solo Android).",
])
p("Importante para anuncios iOS: Misiva y Miravista no tendrán App Store. Lunera, MyPass, ReformaPRO y Monexa sí van por TestFlight → App Store.")

APPS = [
    {
        "name": "Misiva",
        "tagline": "Programa mensajes de WhatsApp con ayuda de IA",
        "importance": "Alta para adquisición Android y retención diaria. Diferencial claro (mensajes a familia + Gemini en el móvil). Canal único Android — no diluir el mensaje con iPhone.",
        "problem": "La gente olvida o pospone mensajes importantes (padres, pareja) y escribir el mismo texto cada día cansa.",
        "for_who": "Personas que quieren cariño o recordatorios automáticos por WhatsApp sin depender de un servidor propio.",
        "features": [
            "Programaciones diarias y recurrentes (contacto, hora, días).",
            "Generación de texto con Gemini (API key del usuario).",
            "Envío en el dispositivo (servicio de accesibilidad + alarmas).",
            "Confirmación opcional antes de enviar.",
            "Checklist y diagnóstico para Xiaomi / envío con app cerrada.",
            "Historial de envíos.",
        ],
        "platforms": "Android APK (landing) y AAB hacia Google Play. Sin iOS.",
        "cta": "Programa tu primer WhatsApp con IA — descárgala en Android.",
        "url": "https://cristianoqa.github.io/descargar-misiva.html",
        "version": "2.0.89",
    },
    {
        "name": "MyPass",
        "tagline": "Tu bóveda de contraseñas, tu clave",
        "importance": "Producto de confianza y privacidad. Ideal para anuncios de seguridad y TestFlight iOS. Sync cloud opcional sin paywall obligatorio del núcleo.",
        "problem": "Los gestores cloud tradicionales no dan control total; hace falta una bóveda local-first con recuperación clara.",
        "for_who": "Usuarios que priorizan privacidad en Android e iPhone (TestFlight).",
        "features": [
            "Bóveda cifrada en el dispositivo + contraseña maestra.",
            "Recovery Key y backup cifrado.",
            "Autofill Android y códigos TOTP.",
            "Biometría opcional.",
            "Sincronización en la nube opcional (datos cifrados).",
            "Centro de ayuda, trust center y tour rejugable.",
        ],
        "platforms": "Android APK · TestFlight iOS · camino a Google Play.",
        "cta": "Tu bóveda, tu clave: prueba MyPass en Android o TestFlight.",
        "url": "https://cristianoqa.github.io/descargar-mypass.html",
        "version": "1.1.8",
    },
    {
        "name": "ReformaPRO",
        "tagline": "Presupuestos de reformas con margen visible",
        "importance": "B2B / autónomos: mayor ticket percibido y uso profesional. La PWA permite probar sin instalar; iOS TestFlight abre canal Apple.",
        "problem": "Autónomos y pymes de reformas pierden margen y tiempo con presupuestos sin catálogo, sin firma y sin control de obra.",
        "for_who": "Autónomos, interioristas y pequeñas constructoras.",
        "features": [
            "Catálogo de materiales/servicios y plantillas.",
            "Presupuestos con IVA, IRPF y margen visible.",
            "Modo cliente (oculta costes internos).",
            "PDF, email y firma digital del cliente.",
            "Dictado por voz a partidas.",
            "Estados: envío → pedido → obra → calendario / obra de hoy.",
            "Equipo multi-usuario e idiomas ES/EN/PT.",
        ],
        "platforms": "PWA (Vercel) · APK Capacitor · TestFlight iOS.",
        "cta": "Abre la web y crea tu primer presupuesto con margen en minutos.",
        "url": "https://reformapro-web-seven.vercel.app",
        "version": "0.2.0",
    },
    {
        "name": "Monexa",
        "tagline": "Gastos e ingresos compartidos en grupo",
        "importance": "Uso diario en pareja/familia. Buen gancho de retención y TestFlight. Diferencial: grupo + roles + OCR/voz.",
        "problem": "Las parejas y grupos no ven juntos gastos, presupuestos y suscripciones en un solo sitio.",
        "for_who": "Hogares y grupos que gestionan dinero juntos.",
        "features": [
            "Grupos con roles (owner / admin / miembro).",
            "Movimientos, categorías, presupuestos, suscripciones y metas.",
            "Insights y revisión semanal / cierre de mes.",
            "OCR de tickets y gasto por voz.",
            "Importación CSV/JSON y modo offline.",
            "Ocultar saldos con biometría.",
        ],
        "platforms": "Android APK · TestFlight iOS · camino a Google Play.",
        "cta": "Controla los gastos del grupo: descarga Monexa y crea tu primer movimiento.",
        "url": "https://cristianoqa.github.io/descargar-monexa.html",
        "version": "1.0.11",
    },
    {
        "name": "Lunera",
        "tagline": "Ciclo y salud femenina con privacidad local",
        "importance": "FemTech con fuerte historia de privacidad. Tres canales (APK, PWA, TestFlight) facilitan difusión. Contenido educativo + tracking.",
        "problem": "Hace falta seguir el ciclo y etapas de vida con orientación educativa y datos que se quedan en el dispositivo.",
        "for_who": "Mujeres que trackean ciclo, anticoncepción, embarazo/postparto o menopausia.",
        "features": [
            "Predicción de fases y registro diario (síntomas, hábitos).",
            "Modos por etapa de vida y objetivos.",
            "Alertas educativas y chat IA local con límites clínicos.",
            "Sección Aprende (enciclopedia).",
            "Informe PDF, backup/sync cifrado.",
            "Modo invitado sin cuenta.",
        ],
        "platforms": "Android APK · PWA · TestFlight iOS · camino a Google Play.",
        "cta": "Conoce tu ciclo en privado: Android, PWA o TestFlight.",
        "url": "https://cristianoqa.github.io/lunera/",
        "version": "1.1.11",
    },
    {
        "name": "Miravista",
        "tagline": "Espejo de pantalla del móvil a la Smart TV",
        "importance": "Utilidad doméstica rápida de entender. Solo Android (límites iOS). Anuncio visual: QR + TV.",
        "problem": "Quieres ver el móvil en la TV sin Chromecast, sin cuenta y sin apps raras en el televisor.",
        "for_who": "Usuarios Android en casa/oficina con TV en la misma Wi‑Fi.",
        "features": [
            "Espejo Wi‑Fi local (MJPEG) al navegador de la TV.",
            "QR + PIN de sesión (máximo 2 visores).",
            "Calidades 480p / 720p / 1080p.",
            "Sin nube, sin anuncios, sin cuenta obligatoria.",
            "Aviso DRM (Netflix etc. no se espejan).",
        ],
        "platforms": "Solo Android APK/AAB. Sin iOS.",
        "cta": "Misma Wi‑Fi, escanea el QR: espejo en tu TV en segundos.",
        "url": "https://cristianoqa.github.io/descargar-miravista.html",
        "version": "1.1.2",
    },
]

h1("2. Fichas por app")
for app in APPS:
    h2(app["name"])
    p(app["tagline"], bold=True)
    h3("Importancia")
    p(app["importance"])
    h3("Problema que resuelve")
    p(app["problem"])
    h3("Para quién")
    p(app["for_who"])
    h3("Funcionalidades clave")
    bullets(app["features"])
    h3("Plataformas y versión")
    p(f"{app['platforms']} · Versión de referencia: {app['version']}")
    h3("CTA sugerida para anuncio")
    p(app["cta"])
    h3("Enlace de descarga / prueba")
    p(app["url"])

h1("3. Mensajes listos para divulgar")
h3("WhatsApp corto — Misiva nueva")
p("¡Misiva nueva! Programa WhatsApp + IA (Gemini) en tu Android.\nDescarga: https://cristianoqa.github.io/descargar-misiva.html\nVersión 2.0.89 — permite «orígenes desconocidos» solo para instalar.")

h3("WhatsApp — catálogo completo")
p(
    "He actualizado mis apps (prueba):\n"
    "• Misiva 2.0.89 (solo Android) → https://cristianoqa.github.io/descargar-misiva.html\n"
    "• Catálogo (MyPass, Monexa, ReformaPRO, Lunera, Miravista) → https://cristianoqa.github.io/\n\n"
    "Misiva y Miravista = solo Android.\n"
    "Lunera, MyPass, ReformaPRO y Monexa también en TestFlight (iPhone, con invitación)."
)

h1("4. Notas para creatividades")
bullets([
    "Un anuncio = una app = un beneficio. No mezclar iOS y solo-Android en el mismo creativo.",
    "Misiva: enfocar familia + WhatsApp + «en tu móvil».",
    "MyPass / Lunera: enfocar privacidad y «tus datos en el dispositivo».",
    "ReformaPRO: enfocar margen / profesional / prueba web inmediata.",
    "Monexa: enfocar pareja/grupo y «primer movimiento en 1 minuto».",
    "Miravista: demo visual QR → TV; dejar claro misma Wi‑Fi y límites DRM.",
    "APK: siempre mencionar activar orígenes desconocidos solo para esa app.",
])

h1("5. Mejoras de ayuda aplicadas (sep 2026)")
bullets([
    "Lunera — FAQ de producto (Ajustes → FAQ), separada de la pestaña Aprende: predicción, registro, invitado vs cuenta, backup, notificaciones, Pro, etapa de vida.",
    "Miravista — Ayuda ampliada: checklist + FAQ (misma Wi‑Fi, Cast vs QR, DRM, firewall/VPN, visores, audio, permisos).",
    "ReformaPRO — FAQ «Ciclo de obra» (presupuesto → firma → pedido → obra/calendario) + tarjeta «Primeros pasos» en el panel (dismissible).",
    "Monexa — Onboarding en Inicio con grupo/roles + FAQ crear/unir/invitar/roles.",
    "Misiva — Banner persistente si falta checklist Estado; FAQ Xiaomi/OEM y «Probar mensaje» vs alarma real; pie Contacto + Privacidad.",
    "MyPass — FAQ Autocompletado iOS (copia/pega + roadmap Credential Provider) + backup gratis vs sync nube opcional.",
])
p(
    "Estas mejoras reducen tickets de soporte y abandono en el primer uso. "
    "Úsalas también en creatividades («cómo empezar en 3 pasos»)."
)

doc.save(OUT)
print(str(OUT))
